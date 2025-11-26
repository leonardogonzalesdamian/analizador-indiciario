from docx import Document
from docx.shared import Pt

def generar_informe_word(buffer, criterios, ici_global, incongruencias):

    doc = Document()

    # Título
    doc.add_heading("Informe de Auditoría Indiciaria (ICI) – Versión 3", level=1)
    doc.add_paragraph(f"Índice ICI Global: {ici_global:.2f}")

    # Tabla criterios
    doc.add_heading("Resultados por criterio C1–C7", level=2)
    table = doc.add_table(rows=1, cols=2)
    hdr = table.rows[0].cells
    hdr[0].text = "Criterio"
    hdr[1].text = "Puntaje"

    for c, v in criterios.items():
        row = table.add_row().cells
        row[0].text = c
        row[1].text = str(v)

    # Incongruencias
    doc.add_heading("Incongruencias detectadas (Reglas 1–9)", level=2)

    if not incongruencias:
        doc.add_paragraph("No se detectaron incongruencias relevantes.")
    else:
        for i, inc in enumerate(incongruencias, start=1):
            doc.add_heading(f"{i}. {inc['tipo']}", level=3)
            doc.add_paragraph(f"Párrafos: {', '.join(str(n) for n in inc['parrafos'])}")
            doc.add_paragraph(inc["detalle"])

            if inc.get("extractos"):
                for ex in inc["extractos"]:
                    p = doc.add_paragraph()
                    run = p.add_run(ex)
                    run.italic = True
                    run.font.size = Pt(10)

    doc.save(buffer)
    buffer.seek(0)
